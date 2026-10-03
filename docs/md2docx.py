# -*- coding: utf-8 -*-
"""把 writeup Markdown 转成 Word（简化解析：标题/表格/代码块/引用/列表/正文）
用法: python md2docx.py [src.md] [dst.docx]"""
import sys, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn

SRC = sys.argv[1] if len(sys.argv) > 1 else r'D:\codes\ctf\docs\2026-09-21_CTF-KCTF-Windows-Reverse-4题-writeup.md'
DST = sys.argv[2] if len(sys.argv) > 2 else SRC.replace('.md', '.docx')

doc = Document()
# 中文字体
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x30, 0x30, 0x30)
    p.paragraph_format.left_indent = Inches(0.25)

def add_table(rows):
    cells = [r.strip().strip('|').split('|') for r in rows]
    cells = [[c.strip() for c in row] for row in cells]
    # 去掉分隔行
    cells = [r for r in cells if not all(re.fullmatch(r':?-+:?', c) for c in r)]
    if not cells: return
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = 'Light Grid Accent 1'
    for i, row in enumerate(cells):
        for j, c in enumerate(row):
            if j < len(t.rows[i].cells):
                t.rows[i].cells[j].text = re.sub(r'\*{2}(.+?)\*{2}', r'\1', re.sub(r'`(.+?)`', r'\1', c))

lines = open(SRC, encoding='utf-8').read().splitlines()
i = 0
in_code = False
code_buf = []
while i < len(lines):
    ln = lines[i]
    if ln.startswith('```'):
        if in_code:
            add_code('\n'.join(code_buf)); code_buf = []
            in_code = False
        else:
            in_code = True
        i += 1; continue
    if in_code:
        code_buf.append(ln); i += 1; continue
    # 表格
    if ln.startswith('|') and i + 1 < len(lines) and lines[i+1].startswith('|'):
        tbl = []
        while i < len(lines) and lines[i].startswith('|'):
            tbl.append(lines[i]); i += 1
        add_table(tbl); continue
    # 标题
    m = re.match(r'^(#{1,4})\s+(.*)', ln)
    if m:
        level = len(m.group(1))
        doc.add_heading(m.group(2).replace('**', ''), level)
        i += 1; continue
    # 引用
    if ln.startswith('> '):
        p = doc.add_paragraph(ln[2:].replace('**', ''))
        p.paragraph_format.left_indent = Inches(0.3)
        for r in p.runs: r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        i += 1; continue
    # 分隔线
    if ln.strip() == '---':
        i += 1; continue
    # 列表
    m = re.match(r'^(\d+)\.\s+(.*)', ln)
    if m:
        doc.add_paragraph(re.sub(r'\*{2}(.+?)\*{2}', r'\1', re.sub(r'`(.+?)`', r'\1', m.group(2))), style='List Number')
        i += 1; continue
    if ln.startswith('- '):
        doc.add_paragraph(re.sub(r'\*{2}(.+?)\*{2}', r'\1', re.sub(r'`(.+?)`', r'\1', ln[2:])), style='List Bullet')
        i += 1; continue
    # 空行
    if not ln.strip():
        i += 1; continue
    # 正文
    doc.add_paragraph(re.sub(r'\*{2}(.+?)\*{2}', r'\1', re.sub(r'`(.+?)`', r'\1', ln)))
    i += 1

doc.save(DST)
print('saved:', DST)
